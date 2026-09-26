import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from tasktracker.services.export import XLSX_MEDIA_TYPE

WEEK = {"starts_at": "2026-09-28T00:00", "ends_at": "2026-10-05T00:00"}


@pytest.fixture
def week(client, make_folder, make_task, make_block):
    """Work and Personal time, plus an inbox appointment, in the week of 28 Sep 2026."""
    work = make_folder("Work", "#2a78d6")
    home = make_folder("Personal", "#eb6834")
    slides = make_task("Slides", folder_id=work["id"])
    groceries = make_task("Groceries", folder_id=home["id"])
    make_block("2026-09-28T09:00", "2026-09-28T10:30", task_id=slides["id"], kind="tracked")
    make_block("2026-09-29T14:00", "2026-09-29T15:00", task_id=slides["id"])
    make_block(
        "2026-09-30T18:00", "2026-09-30T18:45", task_id=groceries["id"], kind="tracked",
        notes="market",
    )  # fmt: skip
    make_block("2026-10-01T12:00", "2026-10-01T12:30", title="Dentist")
    return {"work": work["id"], "home": home["id"], "slides": slides["id"]}


def _report(client, **filters):
    response = client.get("/api/reports/time", params={**WEEK, **filters})
    assert response.status_code == 200, response.text
    return response.json()


def test_report_for_one_folder(client, week):
    report = _report(client, folder_id=week["work"])
    assert (report["tracked_minutes"], report["planned_minutes"]) == (90, 60)
    assert [f["name"] for f in report["folders"]] == ["Work"]
    assert [t["title"] for t in report["tasks"]] == ["Slides"]
    assert sum(d["tracked_minutes"] for d in report["days"]) == 90

    inbox = _report(client, inbox=True)
    assert [t["title"] for t in inbox["tasks"]] == ["Dentist"]
    assert (inbox["tracked_minutes"], inbox["planned_minutes"]) == (0, 30)


def test_completed_tasks_count_per_folder(client, week, make_task):
    done_at = client.patch(f"/api/tasks/{week['slides']}", json={"status": "done"}).json()
    other = make_task("Inbox chore")
    client.patch(f"/api/tasks/{other['id']}", json={"status": "done"})
    day = datetime.fromisoformat(done_at["completed_at"]).date().isoformat()
    around = {"starts_at": f"{day}T00:00", "ends_at": f"{day}T23:59"}

    def completed(**filters):
        return client.get("/api/reports/time", params={**around, **filters}).json()

    assert completed()["completed_tasks"] == 2
    assert completed(folder_id=week["work"])["completed_tasks"] == 1
    assert completed(folder_id=week["home"])["completed_tasks"] == 0
    assert completed(inbox=True)["completed_tasks"] == 1


def test_unknown_folder_is_404(client):
    assert client.get("/api/reports/time", params={**WEEK, "folder_id": 999}).status_code == 404
    assert (
        client.get("/api/reports/time.xlsx", params={**WEEK, "folder_id": 999}).status_code == 404
    )


def _workbook(client, **filters):
    response = client.get("/api/reports/time.xlsx", params={**WEEK, **filters})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    return response, load_workbook(io.BytesIO(response.content))


def _rows(sheet, first_row=1):
    return [list(r) for r in sheet.iter_rows(min_row=first_row, values_only=True)]


def test_export_everything(client, week):
    response, book = _workbook(client)
    assert (
        'filename="time-report_2026-09-28_2026-10-04.xlsx"'
        in response.headers["content-disposition"]
    )
    assert book.sheetnames == ["Summary", "By task", "By day", "Time entries"]

    summary = {r[0]: r[1] for r in _rows(book["Summary"]) if r[0]}
    assert summary["Project"] == "All projects"
    assert summary["From"].date() == date(2026, 9, 28)
    assert summary["To"].date() == date(2026, 10, 4)
    assert summary["Time spent (h)"] == pytest.approx(2.25)
    assert summary["Planned (h)"] == pytest.approx(1.5)

    by_task = _rows(book["By task"])
    assert by_task[0] == ["Task or appointment", "Folder", "Status", "Spent (h)", "Planned (h)"]
    assert [r[:3] for r in by_task[1:4]] == [
        ["Slides", "Work", "To do"],
        ["Groceries", "Personal", "To do"],
        ["Dentist", "Inbox", "Appointment"],
    ]
    assert by_task[-1][0] == "Total"

    by_day = _rows(book["By day"])
    assert len(by_day) == 1 + 7 + 1  # header, seven days, total

    entries = _rows(book["Time entries"])
    assert entries[0][:4] == ["Date", "Start", "End", "Task or appointment"]
    assert [(r[3], r[4], r[5], r[6]) for r in entries[1:5]] == [
        ("Slides", "Work", 1.5, None),
        ("Slides", "Work", None, 1.0),
        ("Groceries", "Personal", 0.75, None),
        ("Dentist", "Inbox", None, 0.5),
    ]
    assert entries[3][7] == "market"


def test_export_one_folder(client, week):
    response, book = _workbook(client, folder_id=week["work"])
    assert "_work.xlsx" in response.headers["content-disposition"]
    summary = {r[0]: r[1] for r in _rows(book["Summary"]) if r[0]}
    assert summary["Project"] == "Work"
    assert summary["Time spent (h)"] == pytest.approx(1.5)
    entries = _rows(book["Time entries"])
    assert {r[4] for r in entries[1:-1]} == {"Work"}


def test_export_an_empty_range(client):
    _, book = _workbook(client)
    assert _rows(book["By task"])[-1][0] == "Total"
